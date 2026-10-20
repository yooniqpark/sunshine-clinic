"""완성된 원고를 워드(.docx) 파일로 만듭니다. 구글 문서·한글(HWP)에서도 열 수 있습니다."""

import io
import re

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

FONT = "맑은 고딕"
KIND_LABEL = {"youtube": "유튜브 대본", "blog": "블로그"}
NAME = {"claude": "Claude", "gpt": "GPT"}


def _set_font(style, size=None):
    style.font.name = FONT
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        style.font.size = Pt(size)


def _add_inline(paragraph, text):
    """**굵게** 표시를 실제 굵은 글씨로 바꿔 넣습니다."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = i % 2 == 1


def _add_markdown(doc, text):
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2).strip(), level=len(m.group(1)))
        elif re.match(r"^\s*[-*•]\s+", line):
            _add_inline(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*•]\s+", "", line))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            _add_inline(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", line))
        else:
            _add_inline(doc.add_paragraph(), line)


def build_docx(draft: str, summary: dict | None = None) -> bytes:
    """원고를 워드 파일로 만듭니다. summary 가 있으면 뒤쪽에 작업 기록 요약을 붙입니다."""
    doc = Document()
    _set_font(doc.styles["Normal"], 11)
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title", "List Bullet", "List Number"):
        _set_font(doc.styles[name])

    _add_markdown(doc, draft)

    if summary:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading("작업 기록", level=1)
        note = doc.add_paragraph()
        note.add_run("이 페이지는 참고용입니다. 글을 올릴 때는 앞쪽 원고만 사용하세요.").font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for k, v in [
            ("글 종류", KIND_LABEL.get(summary.get("kind"), "")),
            ("최종 정리 / 최종 검토" if summary.get("mode") == "debate" else "작가 / 편집자",
             f"{NAME.get(summary.get('writer'), '')} / {NAME.get(summary.get('editor'), '')}"),
            ("점수 변화", summary.get("trajectory") or " → ".join(str(s) for s in summary.get("scores", []))),
            ("종료 이유", summary.get("stop_reason", "")),
        ]:
            row = table.add_row().cells
            row[0].text, row[1].text = k, v

        fb = summary.get("final_feedback") or {}
        if fb.get("fact_check"):
            doc.add_heading("⚠️ 올리기 전에 직접 확인할 것", level=2)
            for s in fb["fact_check"]:
                doc.add_paragraph(s, style="List Bullet")
        if fb.get("strengths"):
            doc.add_heading("편집자가 꼽은 잘된 점", level=2)
            for s in fb["strengths"]:
                doc.add_paragraph(s, style="List Bullet")
        if fb.get("fixes"):
            doc.add_heading("남은 개선 제안", level=2)
            for f in fb["fixes"]:
                doc.add_paragraph(f"{f.get('where')}: {f.get('how')} ({f.get('why')})", style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def read_docx_text(data: bytes) -> str:
    """업로드한 워드 파일에서 글자만 뽑아냅니다."""
    doc = Document(io.BytesIO(data))
    lines = []
    for p in doc.paragraphs:
        style = (p.style.name or "") if p.style is not None else ""
        if style.startswith("Heading") and style[-1].isdigit():
            lines.append("#" * int(style[-1]) + " " + p.text)
        else:
            lines.append(p.text)
    return "\n".join(lines).strip()
