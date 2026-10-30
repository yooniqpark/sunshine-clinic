#!/usr/bin/env python3
"""베가스 '오늘' 팝업에서 읽은 rows.jsonl 을 환자별 표 + 합계로 정리한다."""
import argparse
import csv
import json
import sys

FIELDS = [
    ("total", "총진료비"),
    ("claim", "청구액"),
    ("copay", "본인부담금"),
    ("nonins_tax", "비급여(과세)"),
    ("nonins_notax", "비급여(비과세)"),
    ("discount", "할인"),
]
INFO = [("no", "순번"), ("patient", "환자"), ("chart_no", "차트번호"), ("time", "시간")]


def to_int(value):
    if value in (None, "", "-"):
        return 0
    if isinstance(value, (int, float)):
        return abs(int(value))
    return abs(int(str(value).replace(",", "").replace("원", "").strip() or 0))


def load(path):
    rows = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            for key, _ in FIELDS:
                row[key] = to_int(row.get(key))
            gap = row["total"] - (row["claim"] + row["copay"] + row["nonins_tax"] + row["nonins_notax"])
            if gap and "검산" not in row.get("note", ""):
                row["note"] = (row.get("note", "") + f" 검산차이 {gap:+,}").strip()
            rows.append(row)
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("jsonl")
    parser.add_argument("--out", default="수납정리")
    args = parser.parse_args()

    rows = load(args.jsonl)
    totals = {key: sum(r[key] for r in rows) for key, _ in FIELDS}
    header = [label for _, label in INFO] + [label for _, label in FIELDS] + ["비고"]
    table = [[r.get(k, "") for k, _ in INFO] + [r[k] for k, _ in FIELDS] + [r.get("note", "")] for r in rows]
    total_row = ["합계", f"{len(rows)}명", "", ""] + [totals[k] for k, _ in FIELDS] + [""]

    with open(f"{args.out}.csv", "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(table)
        writer.writerow(total_row)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "수납정리"
        ws.append(header)
        for r in table:
            ws.append(r)
        ws.append(total_row)
        for cell in ws[1] + ws[ws.max_row]:
            cell.font = Font(bold=True)
        for col in range(len(INFO) + 1, len(INFO) + len(FIELDS) + 1):
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                row[0].number_format = "#,##0"
        wb.save(f"{args.out}.xlsx")
        print(f"저장: {args.out}.csv, {args.out}.xlsx")
    except ImportError:
        print(f"저장: {args.out}.csv (openpyxl 없음 → xlsx 생략)")

    print(f"\n환자 수: {len(rows)}명")
    for key, label in FIELDS:
        print(f"{label:<12} {totals[key]:>14,}원")
    flagged = [r for r in rows if r.get("note")]
    if flagged:
        print("\n확인 필요:")
        for r in flagged:
            print(f"  {r.get('no', '')} {r.get('patient', '')}: {r['note']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
